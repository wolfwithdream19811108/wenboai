#!/usr/bin/env python3
"""Batch-test create_for_test in server-enforced dry-run mode.

The source workbook is opened read-only. Every non-empty cell in column C,
starting at row 2 of every selected sheet, is submitted as the final user
message. The script refuses to start unless the endpoint first proves that it
supports dry-run without calling image generation.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import sys
import time
import urllib.error
import urllib.request
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple


DEFAULT_URL = "https://wenboai.cn/api/ai/imagegenerate/create_for_test"
DEFAULT_API_KEY_ENV = "WENBOAI_IMAGE_TEST_API_KEY"
CSV_FIELDS = [
    "case_index",
    "sheet",
    "row",
    "prompt",
    "prompt_hash",
    "cached",
    "http_status",
    "api_status",
    "reason",
    "decision",
    "classifier_enabled",
    "classifier_available",
    "is_sensitive",
    "is_artwork",
    "is_restricted_domain",
    "generation_called",
    "would_generate_in_live_mode",
    "elapsed_ms",
    "attempts",
    "error",
    "tested_at",
    "raw_response",
]


class SafetyError(RuntimeError):
    """Raised when the server cannot prove that image generation was skipped."""


@dataclass(frozen=True)
class TestCase:
    case_index: int
    sheet: str
    row: int
    prompt: str
    prompt_hash: str

    @property
    def key(self) -> Tuple[str, int]:
        return self.sheet, self.row


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="读取 Excel 各 Sheet 的 C 列，并以 dry-run 模式批测 create_for_test。"
    )
    parser.add_argument("workbook", type=Path, help="待测试的 .xlsx 文件")
    parser.add_argument("--url", default=DEFAULT_URL, help="create_for_test 接口地址")
    parser.add_argument(
        "--api-key-env",
        default=DEFAULT_API_KEY_ENV,
        help="保存 Bearer Key 的环境变量名",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("create_for_test_results"),
        help="JSONL 检查点和 CSV 结果目录",
    )
    parser.add_argument("--run-name", help="结果文件名前缀，默认使用 Excel 文件名")
    parser.add_argument("--model", default="jianyuan")
    parser.add_argument("--workers", type=int, default=3, help="并发数，建议 1～5")
    parser.add_argument("--timeout", type=float, default=60.0, help="单次请求超时秒数")
    parser.add_argument("--retries", type=int, default=3, help="网络或 5xx 最大尝试次数")
    parser.add_argument("--limit", type=int, help="只处理前 N 行，用于小规模验证")
    parser.add_argument(
        "--sheets",
        help="只处理指定 Sheet，多个名称用英文逗号分隔",
    )
    parser.add_argument(
        "--no-cache",
        action="store_true",
        help="不复用相同提示词的结果，每一行都单独请求",
    )
    parser.add_argument(
        "--scan-only",
        action="store_true",
        help="只统计待测行，不访问接口",
    )
    return parser.parse_args()


def require_openpyxl():
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit(
            "缺少 openpyxl，请先执行：python -m pip install openpyxl"
        ) from exc
    return load_workbook


def load_cases(workbook_path: Path, selected_sheets: Optional[set]) -> List[TestCase]:
    load_workbook = require_openpyxl()
    workbook = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    cases: List[TestCase] = []
    case_index = 0
    try:
        available_sheets = set(workbook.sheetnames)
        if selected_sheets:
            missing = sorted(selected_sheets - available_sheets)
            if missing:
                raise SystemExit("找不到 Sheet：" + "、".join(missing))

        for sheet in workbook.worksheets:
            if selected_sheets and sheet.title not in selected_sheets:
                continue
            for row_number, values in enumerate(
                sheet.iter_rows(
                    min_row=2,
                    min_col=3,
                    max_col=3,
                    values_only=True,
                ),
                start=2,
            ):
                value = values[0]
                if value is None:
                    continue
                prompt = str(value).strip()
                if not prompt:
                    continue
                case_index += 1
                cases.append(
                    TestCase(
                        case_index=case_index,
                        sheet=sheet.title,
                        row=row_number,
                        prompt=prompt,
                        prompt_hash=hashlib.sha256(prompt.encode("utf-8")).hexdigest(),
                    )
                )
    finally:
        workbook.close()
    return cases


def request_json(
    url: str,
    api_key: str,
    body: Dict[str, Any],
    timeout: float,
    retries: int,
) -> Tuple[int, Dict[str, Any], int, int]:
    encoded = json.dumps(body, ensure_ascii=False).encode("utf-8")
    last_error: Optional[Exception] = None

    for attempt in range(1, retries + 1):
        started = time.perf_counter()
        request = urllib.request.Request(
            url,
            data=encoded,
            method="POST",
            headers={
                "Authorization": "Bearer " + api_key,
                "Content-Type": "application/json; charset=utf-8",
                "Accept": "application/json",
                "User-Agent": "wenboai-dry-run-evaluator/1.0",
            },
        )
        try:
            with urllib.request.urlopen(request, timeout=timeout) as response:
                raw = response.read().decode("utf-8", errors="replace")
                elapsed_ms = int((time.perf_counter() - started) * 1000)
                try:
                    payload = json.loads(raw)
                except json.JSONDecodeError as exc:
                    raise SafetyError("接口未返回 JSON，无法确认 dry-run 安全状态") from exc
                if not isinstance(payload, dict):
                    raise SafetyError("接口返回的 JSON 不是对象，无法确认 dry-run 安全状态")
                return response.status, payload, elapsed_ms, attempt
        except urllib.error.HTTPError as exc:
            elapsed_ms = int((time.perf_counter() - started) * 1000)
            raw = exc.read().decode("utf-8", errors="replace")
            if exc.code >= 500 and attempt < retries:
                last_error = exc
                time.sleep(2 ** (attempt - 1))
                continue
            try:
                payload = json.loads(raw)
            except json.JSONDecodeError:
                payload = {"status": "http_error", "reason": raw[:1000]}
            return exc.code, payload, elapsed_ms, attempt
        except SafetyError:
            raise
        except (urllib.error.URLError, TimeoutError, OSError) as exc:
            last_error = exc
            if attempt < retries:
                time.sleep(2 ** (attempt - 1))
                continue

    raise RuntimeError("请求失败，已达到最大重试次数：" + str(last_error))


def require_dry_run_marker(payload: Dict[str, Any], probe: bool = False) -> Dict[str, Any]:
    data = payload.get("data")
    if not isinstance(data, dict):
        raise SafetyError("响应缺少 data，服务器可能尚未部署 dry-run，已停止测试")
    if data.get("dry_run") is not True or data.get("generation_called") is not False:
        raise SafetyError("服务器未明确证明 generation_called=false，已停止测试")
    if probe and data.get("dry_run_supported") is not True:
        raise SafetyError("服务器未通过 dry-run 安全握手，已停止测试")
    return data


def safety_probe(args: argparse.Namespace, api_key: str) -> None:
    status, payload, _, _ = request_json(
        args.url,
        api_key,
        {"dry_run": True, "dry_run_probe": True},
        args.timeout,
        args.retries,
    )
    if status != 200:
        raise SafetyError("dry-run 安全握手 HTTP 状态码为 %s" % status)
    require_dry_run_marker(payload, probe=True)


def evaluate_prompt(
    args: argparse.Namespace,
    api_key: str,
    prompt: str,
) -> Dict[str, Any]:
    status, payload, elapsed_ms, attempts = request_json(
        args.url,
        api_key,
        {
            "model": args.model,
            "stream": False,
            "dry_run": True,
            "dialogue": [{"role": "user", "content": prompt}],
        },
        args.timeout,
        args.retries,
    )
    data = require_dry_run_marker(payload)
    classification = data.get("classification")
    if not isinstance(classification, dict):
        classification = {}
    return {
        "http_status": status,
        "api_status": payload.get("status", ""),
        "reason": payload.get("reason", ""),
        "decision": data.get("decision", ""),
        "classifier_enabled": classification.get("enabled"),
        "classifier_available": classification.get("available"),
        "is_sensitive": classification.get("is_sensitive"),
        "is_artwork": classification.get("is_artwork"),
        "is_restricted_domain": classification.get("is_restricted_domain"),
        "generation_called": data.get("generation_called"),
        "would_generate_in_live_mode": data.get("would_generate_in_live_mode"),
        "elapsed_ms": elapsed_ms,
        "attempts": attempts,
        "error": "",
        "tested_at": datetime.now(timezone.utc).isoformat(),
        "raw_response": json.dumps(payload, ensure_ascii=False, separators=(",", ":")),
    }


def load_checkpoint(path: Path) -> Dict[Tuple[str, int], Dict[str, Any]]:
    latest: Dict[Tuple[str, int], Dict[str, Any]] = {}
    if not path.exists():
        return latest
    with path.open("r", encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            line = line.strip()
            if not line:
                continue
            try:
                record = json.loads(line)
                latest[(str(record["sheet"]), int(record["row"]))] = record
            except (json.JSONDecodeError, KeyError, TypeError, ValueError):
                print("警告：忽略损坏的检查点第 %d 行" % line_number, file=sys.stderr)
    return latest


def cacheable(record: Dict[str, Any]) -> bool:
    return (
        not record.get("error")
        and record.get("generation_called") is False
        and record.get("decision") in {"allow", "reject"}
    )


def result_for_case(case: TestCase, result: Dict[str, Any], cached: bool) -> Dict[str, Any]:
    row = {
        "case_index": case.case_index,
        "sheet": case.sheet,
        "row": case.row,
        "prompt": case.prompt,
        "prompt_hash": case.prompt_hash,
        "cached": cached,
    }
    row.update(result)
    if cached:
        row["elapsed_ms"] = 0
        row["attempts"] = 0
    return row


def append_jsonl(handle, records: Iterable[Dict[str, Any]]) -> None:
    for record in records:
        handle.write(json.dumps(record, ensure_ascii=False) + "\n")
    handle.flush()


def write_csv(path: Path, records: Iterable[Dict[str, Any]]) -> None:
    temp_path = path.with_suffix(path.suffix + ".tmp")
    with temp_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=CSV_FIELDS, extrasaction="ignore")
        writer.writeheader()
        for record in records:
            writer.writerow(record)
    os.replace(temp_path, path)


def print_scan(cases: List[TestCase]) -> None:
    by_sheet = Counter(case.sheet for case in cases)
    print("有效提示词：%d" % len(cases))
    print("去重提示词：%d" % len({case.prompt_hash for case in cases}))
    print("重复提示词：%d" % (len(cases) - len({case.prompt_hash for case in cases})))
    for sheet, count in by_sheet.items():
        print("  %s: %d" % (sheet, count))


def main() -> int:
    args = parse_args()
    if args.workers < 1 or args.workers > 16:
        raise SystemExit("--workers 必须在 1～16 之间")
    if args.retries < 1 or args.retries > 10:
        raise SystemExit("--retries 必须在 1～10 之间")
    if not args.workbook.is_file():
        raise SystemExit("Excel 文件不存在：%s" % args.workbook)

    selected_sheets = None
    if args.sheets:
        selected_sheets = {name.strip() for name in args.sheets.split(",") if name.strip()}
    cases = load_cases(args.workbook, selected_sheets)
    if args.limit is not None:
        if args.limit < 1:
            raise SystemExit("--limit 必须大于 0")
        cases = cases[: args.limit]
    print_scan(cases)
    if args.scan_only:
        return 0
    if not cases:
        raise SystemExit("没有找到可测试的提示词")

    api_key = os.environ.get(args.api_key_env, "").strip()
    if not api_key:
        raise SystemExit("环境变量 %s 未设置" % args.api_key_env)

    print("正在执行 dry-run 安全握手……")
    safety_probe(args, api_key)
    print("安全握手通过：服务器确认 generation_called=false")

    args.output_dir.mkdir(parents=True, exist_ok=True)
    run_name = args.run_name or args.workbook.stem
    jsonl_path = args.output_dir / (run_name + "_dry_run.jsonl")
    csv_path = args.output_dir / (run_name + "_dry_run.csv")
    checkpoint = load_checkpoint(jsonl_path)
    selected_case_keys = {case.key for case in cases}
    completed = {
        key: record
        for key, record in checkpoint.items()
        if key in selected_case_keys
        and not record.get("error")
        and record.get("generation_called") is False
        and record.get("decision") in {"allow", "reject"}
    }

    prompt_cache: Dict[str, Dict[str, Any]] = {}
    if not args.no_cache:
        for record in checkpoint.values():
            if cacheable(record):
                prompt_cache[str(record.get("prompt_hash", ""))] = {
                    field: record.get(field)
                    for field in CSV_FIELDS
                    if field
                    not in {"case_index", "sheet", "row", "prompt", "prompt_hash", "cached"}
                }

    remaining = [case for case in cases if case.key not in completed]
    groups: Dict[str, List[TestCase]] = defaultdict(list)
    for case in remaining:
        group_key = case.prompt_hash if not args.no_cache else "%s:%d" % case.key
        groups[group_key].append(case)

    written = len(completed)
    total = len(cases)
    with jsonl_path.open("a", encoding="utf-8") as checkpoint_handle:
        for group_key in list(groups):
            group = groups[group_key]
            cached_result = prompt_cache.get(group[0].prompt_hash)
            if cached_result is None:
                continue
            rows = [result_for_case(case, cached_result, True) for case in group]
            append_jsonl(checkpoint_handle, rows)
            for row in rows:
                checkpoint[(row["sheet"], row["row"])] = row
            written += len(rows)
            del groups[group_key]

        try:
            with ThreadPoolExecutor(max_workers=args.workers) as executor:
                futures = {
                    executor.submit(evaluate_prompt, args, api_key, group[0].prompt): group
                    for group in groups.values()
                }
                for future in as_completed(futures):
                    group = futures[future]
                    try:
                        result = future.result()
                    except SafetyError:
                        for pending in futures:
                            pending.cancel()
                        raise
                    except Exception as exc:
                        result = {
                            "http_status": "",
                            "api_status": "",
                            "reason": "",
                            "decision": "error",
                            "classifier_enabled": "",
                            "classifier_available": "",
                            "is_sensitive": "",
                            "is_artwork": "",
                            "is_restricted_domain": "",
                            "generation_called": None,
                            "would_generate_in_live_mode": False,
                            "elapsed_ms": "",
                            "attempts": args.retries,
                            "error": str(exc),
                            "tested_at": datetime.now(timezone.utc).isoformat(),
                            "raw_response": "",
                        }

                    rows = [
                        result_for_case(case, result, index > 0)
                        for index, case in enumerate(group)
                    ]
                    append_jsonl(checkpoint_handle, rows)
                    for row in rows:
                        checkpoint[(row["sheet"], row["row"])] = row
                    written += len(rows)
                    if written == total or written % 25 == 0:
                        print("进度：%d/%d" % (written, total))
        except SafetyError as exc:
            print("安全停止：%s" % exc, file=sys.stderr)
            return 2

    final_records = [
        checkpoint[case.key]
        for case in cases
        if case.key in checkpoint
    ]
    write_csv(csv_path, final_records)
    decisions = Counter(str(record.get("decision", "")) for record in final_records)
    errors = sum(
        1
        for record in final_records
        if record.get("error") or record.get("decision") == "error"
    )
    print("完成。allow=%d reject=%d error=%d" % (
        decisions.get("allow", 0),
        decisions.get("reject", 0),
        errors,
    ))
    print("检查点：%s" % jsonl_path.resolve())
    print("CSV 结果：%s" % csv_path.resolve())
    return 0 if errors == 0 else 1


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("\n用户中止，已完成结果保留在 JSONL 检查点中。", file=sys.stderr)
        raise SystemExit(130)
    except SafetyError as exc:
        print("安全停止：%s" % exc, file=sys.stderr)
        raise SystemExit(2)
